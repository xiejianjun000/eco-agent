# -*- coding: utf-8 -*-
"""Word 紧凑探查：段落/表格统计、标题结构、正文样本、页面设置。
用法: python word_probe.py <文件.docx> [正文样本段数，默认5]
"""
import sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')


def main():
    if len(sys.argv) < 2:
        print("用法: python word_probe.py <文件.docx> [正文样本段数]")
        return 1
    path = sys.argv[1]
    try:
        n = int(sys.argv[2]) if len(sys.argv) > 2 else 5
    except ValueError:
        n = 5

    from docx import Document
    doc = Document(path)
    paras = [p for p in doc.paragraphs if p.text.strip()]

    print(f"文件: {path}")
    print(f"段落: {len(doc.paragraphs)} 个（非空 {len(paras)} 个），表格: {len(doc.tables)} 个")
    for i, t in enumerate(doc.tables, 1):
        print(f"- 表格{i}: {len(t.rows)} 行 x {len(t.columns)} 列")

    def is_heading(p):
        name = getattr(p.style, 'name', '') or ''
        low = name.lower()
        return low.startswith('heading') or low.startswith('标题')

    heads = [p.text for p in paras if is_heading(p)]
    if heads:
        print(f"标题结构（{len(heads)} 个）:")
        for h in heads[:30]:
            print(f"  - {h}")
    else:
        print("标题结构: 无标题样式段落")

    body = [p.text for p in paras if not is_heading(p)]
    print(f"正文前 {n} 段样本:")
    for b in body[:n]:
        print(f"  · {b[:80]}")

    for i, sec in enumerate(doc.sections, 1):
        def fmt(v):
            return f"{v.cm:.1f}" if v is not None else "未设"
        print(f"节{i}: 页边距 上{fmt(sec.top_margin)} 下{fmt(sec.bottom_margin)} "
              f"左{fmt(sec.left_margin)} 右{fmt(sec.right_margin)} cm，"
              f"页面 {fmt(sec.page_width)} x {fmt(sec.page_height)} cm")
    return 0


if __name__ == '__main__':
    sys.exit(main())
