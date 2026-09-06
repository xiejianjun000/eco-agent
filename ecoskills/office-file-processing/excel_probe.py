# -*- coding: utf-8 -*-
"""Excel 紧凑探查：Sheet 列表、行/列数、列名、dtype、前 3 行样本、空值比例。
用法: python excel_probe.py <文件.xlsx> [sheet名|sheet序号]
不指定 sheet 时只概览全部 sheet 的维度（用 openpyxl read_only，速度快、内存省）。
"""
import sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')


def main():
    if len(sys.argv) < 2:
        print("用法: python excel_probe.py <文件路径> [sheet名|sheet序号]")
        return 1
    path = sys.argv[1]
    sel = sys.argv[2] if len(sys.argv) > 2 else None

    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    print(f"文件: {path}")
    print(f"Sheet 列表: {wb.sheetnames}")

    if sel is None:
        # 只概览维度，不读内容
        for name in wb.sheetnames:
            ws = wb[name]
            print(f"- {name}: {ws.max_row} 行 x {ws.max_column} 列")
        wb.close()
        print("提示: 加 sheet 名或序号参数可查看列名与样本，如: ... <文件> Sheet1")
        return 0

    # 解析 sheet 选择（支持名称或 1-based 序号）
    if sel.isdigit():
        idx = int(sel) - 1
        if not (0 <= idx < len(wb.sheetnames)):
            print(f"错误: 序号越界，共 {len(wb.sheetnames)} 个 Sheet")
            return 1
        sel = wb.sheetnames[idx]
    if sel not in wb.sheetnames:
        print(f"错误: 未找到 Sheet [{sel}]，可选: {wb.sheetnames}")
        return 1

    ws = wb[sel]
    nrow, ncol = ws.max_row, ws.max_column
    wb.close()

    import pandas as pd
    df = pd.read_excel(path, sheet_name=sel, nrows=3)
    print(f"Sheet [{sel}]: {nrow} 行 x {ncol} 列")
    print(f"列名: {list(df.columns)}")
    print(f"dtype: {dict(df.dtypes.astype(str))}")

    # 行数不多时顺带算空值比例；大文件跳过以免耗时
    if nrow <= 20000:
        full = pd.read_excel(path, sheet_name=sel)
        miss = {k: float(v) for k, v in full.isna().mean().round(3).items()}
        print(f"空值比例: {miss}")

    print("前 3 行样本:")
    print(df.to_string(max_colwidth=20))
    return 0


if __name__ == '__main__':
    sys.exit(main())
