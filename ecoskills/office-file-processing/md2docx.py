# -*- coding: utf-8 -*-
"""
md2docx.py — Markdown → Word(.docx) 转换器（无外部依赖，仅 python-docx）
用法: python md2docx.py 输入.md [输出.docx]（默认输出与输入同名 .docx）
支持: 标题(#~######)、段落、粗体/斜体/行内代码、链接、无序/有序列表、表格、
      代码块(```)、引用(>)、分隔线(---)
依赖: python-docx（skill 环境已装）
"""
import sys
import re
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8', errors='replace')

from docx import Document
from docx.shared import Pt, RGBColor

INLINE_RE = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')


def add_runs(paragraph, text):
    """解析行内样式: **粗体** *斜体* `代码` [文本](链接)"""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
        elif part.startswith('[') and '](' in part:
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', part)
            if m:
                run = paragraph.add_run(m.group(1))
                run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
                run.underline = True
            else:
                paragraph.add_run(part)
        else:
            paragraph.add_run(part)


def is_table_sep(line):
    """表格分隔行：去掉 | : 空格后只剩 '-'（避免把含 '-' 的数据行误判）"""
    s = line.strip()
    if '|' not in s:
        return False
    core = re.sub(r'[\s|:]', '', s)
    return bool(core) and set(core) == {'-'}


def main():
    if len(sys.argv) < 2:
        print('用法: python md2docx.py 输入.md [输出.docx]')
        sys.exit(1)
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix('.docx')
    lines = src.read_text(encoding='utf-8').splitlines()
    n = len(lines)

    doc = Document()
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)

    i = 0
    while i < n:
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        # 代码块
        if stripped.startswith('```'):
            i += 1
            while i < n and not lines[i].strip().startswith('```'):
                p = doc.add_paragraph()
                run = p.add_run(lines[i] if lines[i] else ' ')
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                p.paragraph_format.space_after = Pt(0)
                i += 1
            i += 1  # 跳过闭合 ```
            continue
        # 标题
        m = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)))
            i += 1
            continue
        # 表格：收集连续 | 行
        if stripped.startswith('|'):
            table_rows = []
            while i < n and lines[i].strip().startswith('|'):
                table_rows.append(lines[i].strip())
                i += 1
            rows = [r for r in table_rows if not is_table_sep(r)]
            if rows:
                data = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
                ncol = max(len(r) for r in data)
                t = doc.add_table(rows=len(data), cols=ncol)
                t.style = 'Table Grid'
                for ri, r in enumerate(data):
                    for ci in range(ncol):
                        cell = t.cell(ri, ci)
                        cell.text = ''
                        add_runs(cell.paragraphs[0], r[ci] if ci < len(r) else '')
                        if ri == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
            continue
        # 分隔线
        if re.match(r'^\s*([-*_])\1{2,}\s*$', stripped):
            p = doc.add_paragraph()
            run = p.add_run('─' * 20)
            run.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
            i += 1
            continue
        # 引用
        if stripped.startswith('>'):
            p = doc.add_paragraph()
            add_runs(p, stripped.lstrip('>').strip())
            for run in p.runs:
                run.italic = True
            p.paragraph_format.left_indent = Pt(24)
            i += 1
            continue
        # 无序列表
        m = re.match(r'^\s*[-*+]\s+(.*)$', stripped)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, m.group(1))
            i += 1
            continue
        # 有序列表
        m = re.match(r'^\s*\d+[.、)]\s+(.*)$', stripped)
        if m:
            p = doc.add_paragraph(style='List Number')
            add_runs(p, m.group(1))
            i += 1
            continue
        # 普通段落
        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    doc.save(dst)
    print(f'✅ 已生成 → {dst}（源 {n} 行）')


if __name__ == '__main__':
    main()
