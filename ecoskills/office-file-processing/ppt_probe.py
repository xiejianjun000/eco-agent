# -*- coding: utf-8 -*-
"""PowerPoint 紧凑探查：幻灯片数、页面尺寸、每页标题与形状统计。
用法: python ppt_probe.py <文件.pptx>
"""
import sys

sys.stdout.reconfigure(encoding='utf-8', errors='replace')


def main():
    if len(sys.argv) < 2:
        print("用法: python ppt_probe.py <文件.pptx>")
        return 1
    path = sys.argv[1]

    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Emu

    prs = Presentation(path)
    print(f"文件: {path}")
    print(f"幻灯片: {len(prs.slides)} 张，"
          f"尺寸 {Emu(prs.slide_width).cm:.1f} x {Emu(prs.slide_height).cm:.1f} cm")

    for i, slide in enumerate(prs.slides, 1):
        ntable = sum(1 for s in slide.shapes if s.has_table)
        npic = sum(1 for s in slide.shapes
                   if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
        title = ''
        if slide.shapes.title is not None:
            title = slide.shapes.title.text.strip().replace('\n', ' / ')
        texts = []
        for s in slide.shapes:
            if s.has_text_frame:
                t = s.text_frame.text.strip().replace('\n', ' / ')
                if t:
                    texts.append(t)
        print(f"[{i}] 形状{len(slide.shapes)} 表格{ntable} 图片{npic} "
              f"| 标题: {title[:50] if title else '(无)'}")
        if texts:
            print(f"    首形状文本: {texts[0][:60]}")
    return 0


if __name__ == '__main__':
    sys.exit(main())
